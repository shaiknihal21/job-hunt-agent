import re
from pathlib import Path

from docx import Document
from langchain_core.messages import HumanMessage, SystemMessage

from app.config import RESUMES_DIR, settings
from app.models.profile import Job, Profile, Resume
from app.services.resume.pdf_export import write_tailored_pdf
from app.services.llm import get_groq_llm
from app.services.resume.ats_scorer import score_ats_alignment
from app.services.resume.evidence import build_candidate_evidence
from app.services.resume.fact_verifier import verify_tailored_content
from app.services.resume.gap_analyzer import analyze_gaps, gaps_missing, gaps_to_strengthen, verified_strengthen_terms
from app.services.resume.jd_parser import parse_job_description
from app.services.resume.role_strategy import get_role_strategy
from app.services.resume.tech_terms import normalize_term, term_in_text
from app.services.resume.types import CandidateEvidence, TailorResult, TailoredContent
from app.services.resume.validator import validate_resume


def _get_llm():
    return get_groq_llm(temperature=0.3, max_tokens=8192)


def _format_gap_report(report) -> str:
    lines = []
    for item in report[:25]:
        lines.append(f"- {item.term}: {item.status.value}")
    return "\n".join(lines)


def _build_tailor_prompt(
    profile: Profile,
    job: Job,
    master_text: str,
    parsed_jd,
    evidence,
    gap_report,
    target_role: str,
    strategy: dict,
) -> str:
    strengthen = gaps_to_strengthen(gap_report)
    missing = gaps_missing(gap_report)

    return f"""Tailor this resume for a specific job application.

CRITICAL ZERO-HALLUCINATION RULES:
- NEVER invent employers, projects, certifications, degrees, metrics, or skills
- Only rephrase and reorder verified content from the master resume and profile skills
- You MAY strengthen MATCHED or SUPPORTED_BUT_NOT_EMPHASIZED skills in wording
- Do NOT add MISSING skills: {', '.join(missing[:15]) or 'none'}
- Do NOT invent percentages or performance metrics

Target role family: {target_role}
Role strategy: {strategy['prompt']}
Company: {job.company}
Job title: {job.title}

Required skills from JD: {', '.join(parsed_jd.required_skills[:20])}
Preferred skills from JD: {', '.join(parsed_jd.preferred_skills[:20])}
Responsibilities: {'; '.join(parsed_jd.responsibilities[:8])}
Keywords to emphasize if verified: {', '.join(strengthen[:20])}

Candidate: {profile.name}, {profile.role}
Profile skills: {', '.join(profile.skills or [])}

Gap analysis:
{_format_gap_report(gap_report)}

- Keep output compact: summary ≤120 words, max 12 skills, max 5 experience bullets, max 3 projects

Master resume:
{master_text[:3500]}

Job description:
{(job.description or '')[:2500]}"""


def _sanitize_content(content: TailoredContent, evidence: CandidateEvidence, gap_report) -> TailoredContent:
    cleaned_skills = []
    for skill in content.skills_reordered:
        norm = normalize_term(skill)
        if term_in_text(norm, evidence.verified_text) or any(
            term_in_text(norm, s) or term_in_text(s, norm) for s in evidence.profile_skills
        ):
            cleaned_skills.append(skill)
    if not cleaned_skills:
        cleaned_skills = verified_strengthen_terms(gap_report, evidence) or evidence.verified_terms[:12]
    content.skills_reordered = cleaned_skills
    return content


def _fallback_content(profile: Profile, job: Job, evidence, gap_report, master_text: str = "") -> TailoredContent:
    verified_skills = verified_strengthen_terms(gap_report, evidence) or evidence.verified_terms[:12] or (profile.skills or [])[:10]
    sentences = [s.strip() for s in re.split(r"[.\n]", master_text) if len(s.strip()) > 25]
    bullets = [f"{s}." if not s.endswith(".") else s for s in sentences[:5]]
    if not bullets:
        bullets = ["See master resume for full experience details."]
    return TailoredContent(
        summary=f"{profile.role} with verified experience in {', '.join(verified_skills[:5])}. "
        f"Applying for {job.title} at {job.company}.",
        skills_reordered=verified_skills,
        experience_bullets=bullets,
        projects=[],
        gaps=gaps_missing(gap_report),
    )


def _generate_content(
    profile: Profile,
    job: Job,
    master_text: str,
    parsed_jd,
    evidence,
    gap_report,
    target_role: str,
    strategy: dict,
) -> TailoredContent:
    llm = _get_llm()
    if not llm:
        return _fallback_content(profile, job, evidence, gap_report, master_text)

    structured = llm.with_structured_output(TailoredContent)
    prompt = _build_tailor_prompt(profile, job, master_text, parsed_jd, evidence, gap_report, target_role, strategy)
    try:
        return structured.invoke([
            SystemMessage(content="You are an expert ATS resume writer. Never fabricate experience. Use ATS-friendly plain language. Keep JSON compact."),
            HumanMessage(content=prompt),
        ])
    except Exception:
        return _fallback_content(profile, job, evidence, gap_report, master_text)


def _safe_filename(company: str, role: str, job_id: int) -> str:
    safe_company = "".join(c if c.isalnum() else "_" for c in company)[:30]
    safe_role = "".join(c if c.isalnum() else "_" for c in role)[:30]
    return f"{safe_company}_{safe_role}_Resume_{job_id}.docx"


def _write_tailored_docx(profile: Profile, job: Job, content: TailoredContent) -> str:
    doc = Document()

    doc.add_heading(profile.name, 0)
    contact_parts = [p for p in [profile.email, profile.phone, profile.location] if p]
    if contact_parts:
        doc.add_paragraph(" | ".join(contact_parts))

    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(content.summary)

    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(", ".join(content.skills_reordered))

    doc.add_heading("Professional Experience", level=1)
    for bullet in content.experience_bullets:
        doc.add_paragraph(bullet, style="List Bullet")

    if content.projects:
        doc.add_heading("Projects", level=1)
        for project in content.projects:
            doc.add_paragraph(project, style="List Bullet")

    filename = _safe_filename(job.company, job.title, job.id)
    file_path = RESUMES_DIR / filename
    doc.save(str(file_path))
    return str(file_path)


def tailor_job_resume(profile: Profile, master_resume: Resume, job: Job) -> TailorResult:
    master_text = master_resume.parsed_text or ""
    parsed_jd = parse_job_description(job)
    evidence = build_candidate_evidence(profile, master_resume, parsed_jd)
    gap_report = analyze_gaps(parsed_jd, evidence)
    target_role, strategy = get_role_strategy(job)

    profile_years = None
    if profile.application_defaults:
        profile_years = profile.application_defaults.get("years_experience")

    ats_before = score_ats_alignment(
        parsed_jd, evidence, master_text, job.title, target_role, profile_years
    )

    content = _generate_content(
        profile, job, master_text, parsed_jd, evidence, gap_report, target_role, strategy
    )
    content = _sanitize_content(content, evidence, gap_report)

    tailored_text = " ".join(
        [content.summary, ", ".join(content.skills_reordered), " ".join(content.experience_bullets)]
    )
    ats_after = score_ats_alignment(
        parsed_jd, evidence, tailored_text, job.title, target_role, profile_years
    )

    fact_passed, fact_issues = verify_tailored_content(content, evidence)
    validation_status, validation_errors = validate_resume(
        content, evidence, parsed_jd, gap_report, ats_after, fact_passed, fact_issues, job, profile
    )

    file_path = _write_tailored_docx(profile, job, content)
    contact = " | ".join(p for p in [profile.email, profile.phone, profile.location] if p)
    pdf_path = write_tailored_pdf(profile.name, contact, content, Path(file_path).with_suffix(".pdf"))

    keywords_emphasized = [
        term for term in gaps_to_strengthen(gap_report)
        if term.lower() in tailored_text.lower()
    ][:15]

    changes_summary = [
        f"Target role: {target_role}",
        f"ATS before: {ats_before.overall} → after: {ats_after.overall}",
        f"Skills reordered: {len(content.skills_reordered)}",
        f"Experience bullets: {len(content.experience_bullets)}",
        f"Gaps flagged: {len(gaps_missing(gap_report))}",
    ]

    content.gaps = gaps_missing(gap_report)

    return TailorResult(
        content=content,
        file_path=file_path,
        target_role=target_role,
        gap_report=gap_report,
        ats_before=ats_before,
        ats_after=ats_after,
        keywords_emphasized=keywords_emphasized,
        changes_summary=changes_summary,
        validation_status=validation_status,
        validation_errors=validation_errors,
        fact_check_passed=fact_passed,
        fact_check_issues=fact_issues,
    )
